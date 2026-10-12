import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner padding for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_banner(doc, title, subtitle, target_info):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(title)
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy Dark

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(4)
    run_sub = p2.add_run(subtitle)
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(2)
    p3.paragraph_format.space_after = Pt(16)
    run_meta = p3.add_run(target_info)
    run_meta.font.name = 'Segoe UI'
    run_meta.font.size = Pt(9.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = 'Segoe UI'
    if level == 1:
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left thick border
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"📌 {title}: ")
    run_t.font.name = 'Segoe UI'
    run_t.font.bold = True
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    run_txt = p.add_run(text)
    run_txt.font.name = 'Segoe UI'
    run_txt.font.size = Pt(9.5)
    run_txt.font.italic = True
    run_txt.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Format header row
    hdr_row = tbl.rows[0]
    hdr_row._element.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "1B365D") # Navy dark
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(heading)
        run.font.name = 'Segoe UI'
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format data rows
    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx + 1]
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(val))
            run.font.name = 'Segoe UI'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
            if c_idx == 0:
                run.font.bold = True

    # Apply subtle border to all cells
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._element.get_or_add_tcPr()
            b = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/></w:tcBorders>')
            tcPr.append(b)

    # Set column widths if provided
    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_code_box(doc, code_str, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(2)
        c_run = cp.add_run(f"💻 {caption}")
        c_run.font.name = 'Segoe UI'
        c_run.font.size = Pt(9.5)
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(bold_prefix + ": ")
    r1.font.name = 'Segoe UI'
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    r2 = p.add_run(text)
    r2.font.name = 'Segoe UI'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

print("Helper functions compiled successfully")
