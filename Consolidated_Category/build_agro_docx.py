import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_helpers import (
    set_cell_background, set_cell_margins, add_header_banner,
    add_styled_heading, add_callout, add_styled_table, add_code_box, add_bullet
)

doc = docx.Document()

# Set standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Header Banner
add_header_banner(
    doc,
    title="B2B MODERN RICE MILL & COMMODITY WEB APPLICATION",
    subtitle="Developer Blueprint: Live Mandi Rates Ticker, Bulk RFQ Estimator & Miller Back-Office",
    target_info="Target Leads: Modern Rice Mills, Paddy Processors, Grain Depots, Agro Exporters | 150 Leads Pipeline"
)

add_callout(
    doc,
    title="DEMO PROTOTYPE DIRECTIVE",
    text="This specification details how to build a B2B commodity wholesale web platform for modern rice millers. For client demonstrations, all grain varieties, technical lab specifications (moisture, broken %, grain length), live mandi prices per quintal, and bulk buyer RFQs are powered by realistic dummy data. No complex server infrastructure is required."
)

# 1. Executive Summary & Purpose
add_styled_heading(doc, "1. Executive Summary & Demo Objective", level=1)
p = doc.add_paragraph()
p.add_run(
    "Rice millers across Andhra Pradesh (Kadapa, Proddatur, Mydukur, Badvel) face two major business handicaps: "
    "(1) Heavy dependence on local commission agents and brokers who take 2-4% cuts and delay payments, and "
    "(2) Out-of-state institutional buyers (in Hyderabad, Bangalore, Chennai, Kerala) have no way to verify their daily prices, "
    "milling capacity, or Sortex machine quality online. "
    "Your objective is to build a professional B2B wholesale portal that showcases the mill's daily grain rates, Sortex infrastructure, "
    "and allows bulk buyers to submit metric-ton RFQs (Requests for Quote) directly to the mill owner's phone."
)

# 2. Frontend Architecture & UI Breakdown
add_styled_heading(doc, "2. Frontend Architecture & Page-by-Page UI Breakdown", level=1)
p2 = doc.add_paragraph()
p2.add_run("The frontend must project industrial credibility, scale, and financial transparency. It consists of 3 primary views:")

add_styled_heading(doc, "2.1 Page 1: Public B2B Mill Showcase ('/')", level=2)
add_bullet(doc, "Live Mandi / Grain Rates Ticker", "Sticky top marquee bar showing today's market rates with 24h trends: '🌾 Sona Masoori Raw: ₹4,200/Qtl [▲ +₹50] | Sona Masoori Steam: ₹3,950/Qtl [▼ -₹20] | Swarna 1061: ₹3,200/Qtl | Broken Rice (Sortex): ₹2,450/Qtl | DORB Bran: ₹2,100/Qtl | Rates Ex-Mill Kadapa'.")
add_bullet(doc, "Industrial Hero Section", "Mill Title (e.g. 'Sri Balaji Modern Rice Mill'), Subtitle ('Leading Processor & Bulk Exporter of Premium Sona Masoori Rice since 1998'), 3 Trust Badges: '150 Metric Tons Daily Capacity', 'Buhler 5-Chute Color Sorter', 'FSSAI & ISO 22000 Certified', and primary CTA ['Request Bulk Wholesale Quote'].")
add_bullet(doc, "Variety Catalog & Technical Lab Specs", "Detailed cards for 5 varieties. Each variety features a technical specification table: Grain Length (e.g. 5.4 mm), Moisture Content (< 12%), Broken Grain (< 1.5%), Purity (99.8% Sortex Clean), and Packaging options.")
add_bullet(doc, "Milling Infrastructure & Storage Silos", "Visual cards highlighting: (1) 5,000 MT Raw Paddy Storage Silos, (2) Buhler Optical Trichromatic Sorter, (3) Automated 25kg/50kg High-Speed Bagging Line, (4) Computerized 60-Ton Weighbridge.")
add_bullet(doc, "Logistics & Transport Connectivity", "Map showing daily truck dispatch routes to Hyderabad, Chennai, Bangalore, and proximity to nearest Railway Goods Siding.")

add_styled_heading(doc, "2.2 Page 2: Bulk RFQ (Request for Quote) Generator ('/rfq')", level=2)
p_rfq = doc.add_paragraph()
p_rfq.add_run("Wholesale buyers do not buy in small retail packs; they buy in Metric Tons (MT) or Full Truck Loads (FTL):")

rfq_components = [
    ["1. Select Rice Variety", "Dropdown / Cards: Sona Masoori Raw, Sona Masoori Steam, Swarna 1061, Broken Rice, DORB Rice Bran."],
    ["2. Enter Quantity (Metric Tons)", "Numeric stepper / slider with quick presets: [10 MT] [20 MT - FTL (Full Truck Load)] [50 MT] [100+ MT Multi-Axle]."],
    ["3. Select Packaging Type", "Radio cards: '25kg Bopp Branded Bag', '50kg Jute Bag (Export Grade)', '50kg PP Non-Woven Bag', 'Plain White Unbranded'."],
    ["4. Destination City", "Delivery hub: Hyderabad, Bangalore, Chennai, Vijayawada, Visakhapatnam, Cochin, or Ex-Mill Kadapa."],
    ["5. Dynamic Cost Estimator", "Interactive calculation displaying estimated Ex-Mill Subtotal (Metric Tons × Live Price/MT) + Packaging Charges = Estimated Total in ₹ Lakhs."],
    ["6. Submit & Send via WhatsApp", "Input: Buyer Company Name, Contact Person, Mobile Number, GSTIN (optional). Tapping submit saves the RFQ and opens WhatsApp with a formal quotation inquiry."]
]
add_styled_table(doc, ["RFQ Step / Component", "Interactive Behavior & Logic"], rfq_components, [2.2, 4.3])

add_styled_heading(doc, "2.3 Page 3: Miller Back-Office Operations Portal ('/portal')", level=2)
p_adm = doc.add_paragraph()
p_adm.add_run("This screen proves to the mill managing partner how they gain complete operational control:")
add_bullet(doc, "1-Click Daily Rate Controller", "A clean table showing each grain variety, current rate per quintal, and rate per 25kg bag. The mill owner can type a new price (e.g. change ₹4,200 to ₹4,250) and tap ['Publish Rates'], instantly updating the public website ticker.")
add_bullet(doc, "Inbound Wholesale RFQ Inbox", "Table of buyer inquiries: RFQ ID, Date, Buyer Company, Variety, Quantity (MT), Destination, Estimated Value, and Status Pill ('New Inquiry' [Blue], 'Quoted' [Amber], 'Confirmed / Dispatched' [Green]).")
add_bullet(doc, "Freight & Truck Dispatch Log", "Summary log of trucks leaving the mill: Lorry Number (e.g. AP 04 TX 9821), Driver Phone, Destination, Net Weight, and Status ('In Transit').")

# 3. Backend Requirements & API Specification
add_styled_heading(doc, "3. Backend Architecture & Mock API Specification", level=1)
p_be = doc.add_paragraph()
p_be.add_run(
    "For the demo application, all endpoints can run in Next.js API Routes, Express.js, or local state with localStorage persistence. "
    "Here is the API contract for the developer:"
)

api_endpoints = [
    ["GET", "/api/mill/rates", "Returns array of all rice varieties with today's live rates per quintal/bag and 24h trend."],
    ["PUT", "/api/mill/rates/:id", "Payload: { pricePerQuintal, trend }. Updates rate for a variety from miller back-office."],
    ["GET", "/api/mill/varieties", "Returns detailed technical grain specifications (grain length, moisture, broken %)."],
    ["POST", "/api/mill/rfq", "Payload: { varietyId, quantityMt, packagingType, destinationCity, buyerCompany, phone }. Calculates quote, logs RFQ, and returns RFQ reference number."],
    ["GET", "/api/mill/portal/rfqs", "Returns list of all buyer inquiries for the miller's management dashboard."]
]
add_styled_table(doc, ["Method", "Endpoint Route", "Purpose & Behavior"], api_endpoints, [1.0, 2.5, 3.0])

# 4. WhatsApp Wholesale Inquiry Engine
add_styled_heading(doc, "4. WhatsApp Wholesale Inquiry Integration Engine", level=1)
p_wa = doc.add_paragraph()
p_wa.add_run(
    "In agricultural commodities, transactions are negotiated over phone and WhatsApp. "
    "When a buyer submits an RFQ on the website, this utility compiles a professional quotation inquiry and deep-links to the miller's WhatsApp:"
)

wa_code = """// utils/whatsappRFQ.ts
export function generateWhatsAppRFQUrl(
  millerPhone: string,
  millName: string,
  rfqNumber: string,
  buyerCompany: string,
  contactPerson: string,
  buyerPhone: string,
  varietyName: string,
  quantityMt: number,
  packaging: string,
  destinationCity: string,
  estimatedTotal: number
): string {
  const formattedAmount = (estimatedTotal / 100000).toFixed(2); // Convert to Lakhs

  let msg = `🌾 *WHOLESALE BULK RFQ - ${millName.toUpperCase()}*\\n`;
  msg += `──────────────────────────\\n`;
  msg += `📋 *Inquiry Ref*: *${rfqNumber}*\\n`;
  msg += `🏢 *Buyer*: ${buyerCompany}\\n`;
  msg += `👤 *Contact*: ${contactPerson} (${buyerPhone})\\n`;
  msg += `──────────────────────────\\n`;
  msg += `📦 *COMMODITY DETAILS*:\\n`;
  msg += `• *Variety*: ${varietyName}\\n`;
  msg += `• *Volume*: *${quantityMt} Metric Tons* (~${quantityMt * 40} Bags of 25kg)\\n`;
  msg += `• *Packaging*: ${packaging}\\n`;
  msg += `• *Delivery Hub*: ${destinationCity}\\n`;
  msg += `──────────────────────────\\n`;
  msg += `💰 *Estimated Ex-Mill Value*: *₹${formattedAmount} Lakhs*\\n`;
  msg += `──────────────────────────\\n`;
  msg += `Please confirm truckload availability and best freight quotation. Thank you!`;

  const cleanPhone = millerPhone.replace(/[^0-9]/g, '');
  return `https://wa.me/${cleanPhone}?text=${encodeURIComponent(msg)}`;
}"""
add_code_box(doc, wa_code, "TypeScript: WhatsApp Bulk RFQ URL Generator")

# 5. Database Schema & Data Models
add_styled_heading(doc, "5. Relational Database Schema (SQL DDL)", level=1)
p_sql = doc.add_paragraph()
p_sql.add_run("Use this relational schema for SQLite, PostgreSQL, or local state data modeling:")

sql_code = """-- Rice Mill Profile
CREATE TABLE rice_mills (
    mill_id TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    managing_partner TEXT NOT NULL,
    phone TEXT NOT NULL,
    email TEXT,
    city TEXT NOT NULL,
    daily_milling_capacity_tpd INTEGER DEFAULT 150,
    sorter_tech TEXT DEFAULT 'Buhler 5-Chute Trichromatic Color Sorter',
    fssai_license TEXT NOT NULL
);

-- Grain Varieties & Live Mandi Pricing
CREATE TABLE grain_varieties (
    variety_id TEXT PRIMARY KEY,
    mill_id TEXT REFERENCES rice_mills(mill_id),
    name TEXT NOT NULL,
    grain_length_mm REAL NOT NULL,
    moisture_pct REAL NOT NULL,
    broken_pct REAL NOT NULL,
    price_per_quintal_inr REAL NOT NULL, -- 100 kg
    price_per_25kg_inr REAL NOT NULL,
    trend_24h TEXT DEFAULT '+0',        -- '+50', '-20', '0'
    in_stock_tonnage INTEGER DEFAULT 500,
    last_updated TIMESTAMP DEFAULT CURRENT_TIMESTAMP
);

-- Bulk RFQs (Wholesale Inquiries)
CREATE TABLE bulk_rfqs (
    rfq_id TEXT PRIMARY KEY,
    rfq_number TEXT UNIQUE NOT NULL,
    variety_id TEXT REFERENCES grain_varieties(variety_id),
    buyer_company TEXT NOT NULL,
    contact_person TEXT NOT NULL,
    phone TEXT NOT NULL,
    quantity_metric_tons REAL NOT NULL,
    packaging_type TEXT NOT NULL,
    destination_city TEXT NOT NULL,
    estimated_amount_inr REAL NOT NULL,
    status TEXT DEFAULT 'NEW_INQUIRY', -- 'NEW_INQUIRY', 'QUOTED', 'CONFIRMED'
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
);"""
add_code_box(doc, sql_code, "SQL Schema for Agro Rice Mill App")

# 6. Complete Dummy Data Fixture
add_styled_heading(doc, "6. Ready-to-Use Dummy Data Fixture (JSON)", level=1)
p_fix = doc.add_paragraph()
p_fix.add_run("Copy this exact JSON into `src/data/agro-mock.json` to immediately populate grain varieties and rates:")

json_fixture = """{
  "mill": {
    "name": "Sri Balaji Modern Rice Mill",
    "managing_partner": "B. Balaji Subbaiah",
    "phone": "+91 98480 33202",
    "email": "balajiricemill.kdp@gmail.com",
    "city": "Kadapa",
    "district": "Kadapa (YSR Kadapa)",
    "address": "Industrial Development Area (IDA), Ramanjaneya Puram, Kadapa - 516002",
    "daily_capacity": "150 Metric Tons/Day (~6,000 Bags/Day)",
    "sorter": "Buhler 5-Chute Trichromatic Optical Color Sorter",
    "certifications": [
      "FSSAI Central License #10118023000412",
      "ISO 22000:2018 Food Safety Management",
      "APEDA Registered Exporter"
    ],
    "silo_capacity": "5,000 Metric Tons Paddy Storage"
  },
  "grain_rates": [
    {
      "id": "grn-01",
      "name": "Sona Masoori Raw Rice (Super Fine Grade)",
      "grain_length": "5.4 mm",
      "moisture": "12.0%",
      "broken": "1.5%",
      "purity": "99.8% Sortex Clean",
      "price_per_quintal": 4200,
      "price_per_25kg": 1050,
      "trend_24h": "+50",
      "stock_mt": 450
    },
    {
      "id": "grn-02",
      "name": "Sona Masoori Steam Rice (Standard Commercial)",
      "grain_length": "5.3 mm",
      "moisture": "12.5%",
      "broken": "2.0%",
      "purity": "99.5% Sortex Clean",
      "price_per_quintal": 3950,
      "price_per_25kg": 990,
      "trend_24h": "-20",
      "stock_mt": 800
    },
    {
      "id": "grn-03",
      "name": "Swarna 1061 Boiled / Parboiled Rice",
      "grain_length": "5.0 mm",
      "moisture": "13.0%",
      "broken": "3.0%",
      "purity": "99.0% Cleaned",
      "price_per_quintal": 3200,
      "price_per_25kg": 800,
      "trend_24h": "0",
      "stock_mt": 1200
    },
    {
      "id": "grn-04",
      "name": "100% Sortex Broken Rice (Nookalu)",
      "grain_length": "2.5 mm",
      "moisture": "11.5%",
      "broken": "100%",
      "purity": "Cleaned Dust-Free",
      "price_per_quintal": 2450,
      "price_per_25kg": 612,
      "trend_24h": "+30",
      "stock_mt": 250
    },
    {
      "id": "grn-05",
      "name": "De-Oiled Rice Bran (DORB) - 16% Protein",
      "grain_length": "Standard Powder",
      "moisture": "10.0%",
      "broken": "N/A",
      "purity": "Pure Feed Grade",
      "price_per_quintal": 2100,
      "price_per_25kg": 525,
      "trend_24h": "+40",
      "stock_mt": 600
    }
  ],
  "recent_rfqs": [
    {
      "rfq_number": "RFQ-9021",
      "buyer_company": "Sri Sai Srinivasa Wholesale Traders",
      "contact_person": "P. Srinivasa Reddy",
      "destination": "Hyderabad, Telangana",
      "variety": "Sona Masoori Raw Rice",
      "quantity_mt": 25.0,
      "packaging": "25kg BOPP Branded Bag",
      "estimated_total_inr": 1050000,
      "status": "QUOTED"
    },
    {
      "rfq_number": "RFQ-9022",
      "buyer_company": "Annapurna Food Corporation",
      "contact_person": "M. Venkatramana",
      "destination": "Chennai Port, Tamil Nadu",
      "variety": "Sona Masoori Steam Rice",
      "quantity_mt": 50.0,
      "packaging": "50kg Jute Export Bag",
      "estimated_total_inr": 1975000,
      "status": "NEW_INQUIRY"
    }
  ]
}"""
add_code_box(doc, json_fixture, "JSON Data Fixture: agro-mock.json")

# 7. Recommended Tech Stack & Folder Structure
add_styled_heading(doc, "7. Recommended Tech Stack & Folder Structure", level=1)
p_stk = doc.add_paragraph()
p_stk.add_run(
    "• Framework: Next.js 14+ (App Router) OR Vite + React 18+\n"
    "• Styling: TailwindCSS (Industrial Navy: `#1E3A8A` + Harvest Amber: `#D97706`)\n"
    "• Icons: Lucide React (TrendingUp, TrendingDown, Truck, ShieldCheck, Factory, FileText, Phone)\n"
    "• State Management: React Context or Zustand for Live Rates and RFQ state."
)

dir_tree = """agro-demo/
├── src/
│   ├── app/ (or pages/)
│   │   ├── page.tsx               # Public Mill Showcase & Spec Sheet
│   │   ├── rfq/page.tsx           # Bulk RFQ Calculator & Submit
│   │   └── portal/page.tsx        # Miller Back-Office Price Controller
│   ├── components/
│   │   ├── LiveRateTicker.tsx     # Sticky top ticker showing live mandi rates
│   │   ├── HeroIndustrial.tsx     # Mill capacity, Sortex machine badges
│   │   ├── VarietySpecCard.tsx    # Technical grain lab specs (moisture, broken %)
│   │   ├── BulkEstimator.tsx      # Metric-ton slider with live cost preview
│   │   └── MillerRateEditor.tsx   # 1-click price updater for mill owner
│   ├── data/
│   │   └── agro-mock.json         # The dummy data fixture provided above
│   └── utils/
│       └── whatsappRFQ.ts         # WhatsApp bulk quote formatter"""
add_code_box(doc, dir_tree, "Clean Repository Structure")

# 8. Step-by-Step Developer Implementation Roadmap
add_styled_heading(doc, "8. Step-by-Step Developer Implementation Roadmap", level=1)
steps = [
    ["Day 1: Industrial Showcase & Ticker", "Build sticky live mandi rate ticker bar with green/red trend badges. Build hero section highlighting 150 TPD capacity and Sortex technology."],
    ["Day 2: Variety Catalog & Spec Tables", "Render 5 grain variety cards with technical lab specs (grain length mm, moisture %, broken %). Add 'Request Quote' CTA."],
    ["Day 3: Bulk RFQ Metric-Ton Calculator", "Build `/rfq` wizard. Implement dynamic formula: Quantity (MT) × Price/MT + Packaging = Estimated Quote in ₹ Lakhs."],
    ["Day 4: WhatsApp Engine & Miller Portal", "Integrate WhatsApp deep-link. Build `/portal` allowing miller to update prices and view inbound RFQs."],
    ["Day 5: Mobile & Tablet Responsiveness", "Verify ticker scrolls smoothly on mobile and tables resize gracefully."]
]
add_styled_table(doc, ["Milestone", "Actionable Deliverables"], steps, [2.0, 4.5])

# 9. Client Demo Pitch Script
add_styled_heading(doc, "9. How to Demonstrate This App to a Rice Mill Owner", level=1)
p_pitch = doc.add_paragraph()
p_pitch.add_run(
    "When presenting to B. Balaji Subbaiah or any rice mill owner from our leads list, follow this 3-minute demonstration flow:\n"
    "1. Show the homepage on a tablet/laptop: 'Balaji garu, this is your mill's digital presence. When out-of-state buyers in Hyderabad or Bangalore open this, the first thing they see is today's live rate ticker and your Buhler Sortex machine credentials.'\n"
    "2. Navigate to the RFQ Calculator: 'A supermarket chain in Chennai wants 25 Metric Tons of Sona Masoori Raw. They slide the volume to 25 MT, choose 25kg BOPP bags, and the system instantly estimates ₹10.50 Lakhs.'\n"
    "3. Tap Submit RFQ: 'The entire truckload inquiry arrives straight to your WhatsApp. You cut out the middleman broker and negotiate directly.'\n"
    "4. Switch to `/portal`: 'And every morning at 9:00 AM, you open your phone, update today's rate from ₹4,200 to ₹4,250, tap Publish, and your prices update across the entire market instantly.'"
)

doc.save("Rice_Mill_and_Agro_Industry_Web_App_Developer_Guide.docx")
print("Saved: Rice_Mill_and_Agro_Industry_Web_App_Developer_Guide.docx")
