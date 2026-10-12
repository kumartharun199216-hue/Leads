import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter, defaultdict

# Load existing clean master dataset without modifying it
source_path = 'Consolidated_Leads_and_Web_Development_Directory.xlsx'
if not os.path.exists(source_path):
    source_path = os.path.join('..', 'Consolidated_Category', 'Consolidated_Leads_and_Web_Development_Directory.xlsx')
source_wb = openpyxl.load_workbook(source_path, data_only=True)
source_ws = source_wb['All Leads (Master Dataset)']

# Micro-classification logic
def get_micro_classification(cat, sub_ind, name):
    cat_low = str(cat or '').lower()
    ind_low = str(sub_ind or '').lower()
    name_low = str(name or '').lower()
    
    if 'healthcare' in cat_low or 'clinic' in cat_low:
        if 'dental' in ind_low or 'dental' in name_low or 'tooth' in ind_low or 'teeth' in ind_low or 'orthodont' in ind_low or 'implant' in ind_low:
            return (
                'Dental Clinics & Implant Centers',
                'Dental Practice & Implant Booking Portal',
                'Procedure menu (RCT, Implants, Braces, Whitening), Morning/Evening slot picker, instant Token ID, WhatsApp booking alert',
                'Slot availability validation (/api/slots), Appointment booking state, Token counter'
            )
        elif any(x in ind_low or x in name_low for x in ['derm', 'skin', 'cosmet', 'tricholog', 'hair', 'aesthetic']):
            return (
                'Dermatology, Skin & Cosmetology Clinics',
                'Aesthetic Dermatology & Skin Care Portal',
                'Treatment catalog (Acne scar, Laser, Chemical peel, HydraFacial), Before/After image slider, Doctor consultation booking, Skincare product store',
                'Consultation slot manager, Treatment inquiry log, WhatsApp appointment trigger'
            )
        elif any(x in ind_low or x in name_low for x in ['pediatr', 'child', 'neonat']):
            return (
                'Pediatrics & Child Care Clinics',
                'Pediatric Care & Baby Immunization Portal',
                'Baby vaccination milestone schedule tracker, Pediatrician slot booking, Growth chart calculator, WhatsApp vaccine alert trigger',
                'Vaccine schedule API, Age-based immunization calculator, Appointment logger'
            )
        elif any(x in ind_low or x in name_low for x in ['eye', 'ophthalm', 'retina', 'cornea', 'cataract', 'glaucoma']):
            return (
                'Ophthalmology & Eye Care Clinics',
                'Eye Care & Optical Storefront Portal',
                'Vision test appointment booking, Cataract/Lasik surgery information guide, Doctor schedule, Spectacle frames & optical catalog',
                'Vision check slot booking, Optical product catalog API, Consultation queue'
            )
        elif any(x in ind_low or x in name_low for x in ['ortho', 'joint', 'bone', 'spine', 'fracture', 'trauma', 'arthritis']):
            return (
                'Orthopedics & Joint Care Clinics',
                'Orthopedic & Joint Pain Clinic Portal',
                'Joint pain self-assessment quiz, Orthopedic surgeon booking, Physiotherapy slot selector, Post-op rehabilitation video guides',
                'Assessment score calculator, Consultation slot engine, WhatsApp booking notifier'
            )
        elif any(x in ind_low or x in name_low for x in ['diabet', 'endocrin', 'thyroid', 'metabolic']):
            return (
                'Diabetology & Endocrinology Centers',
                'Diabetology & Chronic Lifestyle Care Portal',
                'Fasting blood sugar & HbA1c self-logger, Dietician consultation booking, Diabetic foot care clinic scheduler, Teleconsultation CTA',
                'Patient sugar logs mock API, Diet plan downloads, Doctor appointment queue'
            )
        elif any(x in ind_low or x in name_low for x in ['ent', 'ear', 'audiolog', 'hearing', 'speech']):
            return (
                'ENT & Hearing Care Clinics',
                'ENT & Audiology Assessment Portal',
                'Hearing checkup booking, Sinus & allergy consultation scheduler, Ear endoscopy information, Digital hearing aids showcase',
                'Audiogram test booking API, Doctor schedule manager, WhatsApp confirmation'
            )
        elif any(x in ind_low or x in name_low for x in ['homeo', 'ayurv', 'panchakarma', 'herbal', 'holistic']):
            return (
                'Ayurveda & Homeopathy Wellness Centers',
                'Holistic Healing & Natural Wellness Portal',
                'Prakriti / Body constitution quiz, Panchakarma rejuvenation therapy booking, Herbal remedies catalog, Long-term case-taking scheduler',
                'Holistic therapy booking API, Questionnaire response evaluator, Appointment queue'
            )
        else:
            return (
                'Multi-Specialty & General Outpatient Clinics',
                'Multi-Specialty Polyclinic Management Portal',
                'Multi-doctor department filter (Cardio, Gastro, Pulmo, General), Centralized token queue, Diagnostic test booking, Emergency CTA',
                'Multi-department slot engine, Centralized receptionist token board, Patient history API'
            )
            
    elif 'retail' in cat_low or 'kirana' in cat_low:
        if any(x in ind_low or x in name_low for x in ['wholesale', 'depot', 'grain merchant', 'bulk']):
            return (
                'Wholesale Grain Depots & Rice Merchants',
                'B2B Wholesale Grain & Merchant Order Portal',
                'Bulk bag pricing (25kg/50kg), Wholesale tiered discount calculator, Merchant credit khata, GST invoice generator, WhatsApp bulk cart',
                'Tiered pricing API, Bulk cart formatter, Merchant credit ledger manager'
            )
        elif any(x in ind_low or x in name_low for x in ['mini mart', 'superette', 'daily needs', 'dept store', 'fancy', 'daily mart']):
            return (
                'Daily Needs Marts & Superettes',
                'Neighborhood Mini-Mart Fast Order App',
                'Packaged FMCG, Dairy & Cold Beverages, Snacks, Fast autocomplete search, Sticky floating cart, 1-Click WhatsApp delivery order',
                'Product inventory search API, WhatsApp message compiler, Order logger'
            )
        else:
            return (
                'Neighborhood Kirana & General Stores',
                'Hyperlocal Digital Kirana PWA & WhatsApp Store',
                'Daily staples (Rice, Dals, Oils, Spices), 2-column mobile card, inline quantity incrementer (- 1 +), 1-Click WhatsApp order, Free 30-min delivery banner',
                'Product catalog state, Instant price updater, WhatsApp order generator, Digital Khata API'
            )
            
    else: # Agro
        if 'processing' in ind_low or 'food' in ind_low:
            return (
                'Rice Processing & Value-Add Mills',
                'Agro Commodity & By-Product Procurement Portal',
                'Broken Rice (Nookalu) & Rice Bran (DORB) specs, Protein/moisture lab reports, Bulk MT inquiry form, Institutional buyer contracts',
                'Commodity spec sheet API, Bulk RFQ calculator, Contract inquiry manager'
            )
        else:
            return (
                'Modern Rice Mills & Commercial Hullers',
                'B2B Industrial Miller Platform & Live Mandi Ticker',
                'Sticky live mandi rate ticker (Raw/Steam Sona Masoori, Swarna 1061), Buhler Sortex Optical Sorter showcase, Bulk RFQ (Metric Tons) calculator, Freight dispatch log',
                'Live Mandi rate ticker API, 1-click price controller, Bulk RFQ calculation engine, Lorry dispatch tracker'
            )

# Parse all records from source
master_records = []
for r in range(2, source_ws.max_row + 1):
    lead_id = source_ws.cell(r, 1).value
    broad_cat = source_ws.cell(r, 2).value
    bname = source_ws.cell(r, 3).value
    owner = source_ws.cell(r, 4).value
    role = source_ws.cell(r, 5).value
    email = source_ws.cell(r, 6).value
    status = source_ws.cell(r, 7).value
    phone = source_ws.cell(r, 8).value
    city = source_ws.cell(r, 9).value
    district = source_ws.cell(r, 10).value
    state_country = source_ws.cell(r, 11).value
    sub_ind = source_ws.cell(r, 12).value
    gps = source_ws.cell(r, 13).value
    source_info = source_ws.cell(r, 15).value
    source_file = source_ws.cell(r, 16).value
    
    micro_type, app_name, fe_features, be_reqs = get_micro_classification(broad_cat, sub_ind, bname)
    
    master_records.append({
        'lead_id': lead_id,
        'broad_cat': broad_cat,
        'micro_type': micro_type,
        'app_name': app_name,
        'fe_features': fe_features,
        'be_reqs': be_reqs,
        'bname': bname,
        'owner': owner,
        'role': role,
        'email': email,
        'status': status,
        'phone': phone,
        'city': city,
        'district': district,
        'state_country': state_country,
        'sub_ind': sub_ind,
        'gps': gps,
        'source_info': source_info,
        'source_file': source_file
    })

source_wb.close()
print(f'Parsed {len(master_records)} leads from master dataset.')

# Create NEW Workbook
wb_new = openpyxl.Workbook()
wb_new.remove(wb_new.active)

# Color Palette & Styles
navy_dark = '1B365D'
navy_med = '2F5597'
navy_light = 'D9E1F2'
teal_dark = '004C6D'
card_bg = 'F2F4F8'
zebra_bg = 'F9FAFC'
green_text = '1E6B37'
green_bg = 'E2EFDA'
amber_text = '8C5900'
amber_bg = 'FFF2CC'
purple_text = '5E2750'
purple_bg = 'F2E7F0'

font_title = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='595959')
font_section = Font(name='Segoe UI', size=12, bold=True, color=navy_dark)
font_tbl_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_data = Font(name='Segoe UI', size=9, bold=False, color='000000')
font_data_bold = Font(name='Segoe UI', size=9, bold=True, color='000000')
font_card_num = Font(name='Segoe UI', size=18, bold=True, color=navy_dark)
font_card_lbl = Font(name='Segoe UI', size=9, bold=True, color='4A5568')
font_card_sub = Font(name='Segoe UI', size=8, italic=True, color='718096')

fill_navy_header = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type='solid')
fill_sub_header = PatternFill(start_color=navy_med, end_color=navy_med, fill_type='solid')
fill_zebra = PatternFill(start_color=zebra_bg, end_color=zebra_bg, fill_type='solid')
fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
fill_card = PatternFill(start_color=card_bg, end_color=card_bg, fill_type='solid')
fill_green = PatternFill(start_color=green_bg, end_color=green_bg, fill_type='solid')
fill_amber = PatternFill(start_color=amber_bg, end_color=amber_bg, fill_type='solid')

thin_side = Side(border_style='thin', color='D0D5DD')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
double_bottom = Border(top=thin_side, bottom=Side(border_style='double', color=navy_dark))

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_wrap_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

# ==========================================================
# SHEET 1: Micro-Category Analytics & Matrix
# ==========================================================
ws1 = wb_new.create_sheet('Micro Analytics & Matrix')
ws1.views.sheetView[0].showGridLines = True

# Title Banner
ws1.merge_cells('A1:J2')
top_cell = ws1['A1']
top_cell.value = 'MICRO-ACCURATE BUSINESS CLASSIFICATION & WEB APP FEATURE MATRIX'
top_cell.font = font_title
top_cell.fill = fill_navy_header
top_cell.alignment = align_center

ws1['A3'].value = 'Granular Breakdown of 550 Leads into 14 Distinct Commercial Niches with Custom Frontend & Backend Demo Specifications'
ws1['A3'].font = font_subtitle

# KPI Summary Cards (Row 5 - 7)
kpis = [
    ('TOTAL LEADS', '550', '100% Granular Classification', 'B', 'C'),
    ('MICRO BUSINESS NICHES', '14', 'Custom Web App Specs', 'D', 'E'),
    ('DENTAL PRACTICES', '120', 'Smile Portals & Booking', 'F', 'G'),
    ('MEDICAL SPECIALTIES', '80', 'Skin, Child, Eye, Ortho, ENT', 'H', 'I'),
    ('KIRANA & RETAIL STORES', '178', 'WhatsApp Storefront & Khata', 'J', 'K'),
    ('GRAIN DEPOTS & MILLS', '172', 'B2B Portals & Live Tickers', 'L', 'M')
]

for title, val, sub, c1, c2 in kpis:
    ws1.merge_cells(f'{c1}5:{c2}5')
    ws1.merge_cells(f'{c1}6:{c2}6')
    ws1.merge_cells(f'{c1}7:{c2}7')
    
    cell_lbl = ws1[f'{c1}5']
    cell_lbl.value = title
    cell_lbl.font = font_card_lbl
    cell_lbl.alignment = align_center
    cell_lbl.fill = fill_card
    
    cell_val = ws1[f'{c1}6']
    cell_val.value = val
    cell_val.font = font_card_num
    cell_val.alignment = align_center
    cell_val.fill = fill_card
    
    cell_sub = ws1[f'{c1}7']
    cell_sub.value = sub
    cell_sub.font = font_card_sub
    cell_sub.alignment = align_center
    cell_sub.fill = fill_card
    
    for r in range(5, 8):
        for col_letter in [c1, c2]:
            ws1[f'{col_letter}{r}'].border = thin_border

# TABLE 1: Micro-Category Breakdown Table
ws1['A9'].value = '1. Granular Micro-Business Classification & Custom Demo Application Specifications'
ws1['A9'].font = font_section

tbl1_headers = [
    'Primary Sector',
    'Micro-Business Niche',
    'Lead Count',
    '% Share of Total',
    'Tailored Web Application Archetype',
    'Essential Frontend Demo Features',
    'Backend Demo Requirements'
]

for col_idx, h in enumerate(tbl1_headers, start=1):
    cell = ws1.cell(10, col_idx, h)
    cell.font = font_tbl_header
    cell.fill = fill_navy_header
    cell.alignment = align_center if col_idx in [3, 4] else align_left
    cell.border = thin_border

# Aggregate by micro_type
micro_counts = Counter([r['micro_type'] for r in master_records])
# Predefined order with detailed specifications
micro_specs_order = [
    ('Healthcare & Clinics', 'Dental Clinics & Implant Centers', 'Dental Practice & Implant Booking Portal', 'Procedure menu (RCT, Implants, Braces, Whitening), Morning/Evening slot picker, instant Token ID, WhatsApp booking alert', 'Slot availability validation (/api/slots), Appointment booking state, Token counter'),
    ('Healthcare & Clinics', 'Dermatology, Skin & Cosmetology Clinics', 'Aesthetic Dermatology & Skin Care Portal', 'Treatment catalog (Acne scar, Laser, Chemical peel, HydraFacial), Before/After image slider, Doctor consultation booking, Skincare product store', 'Consultation slot manager, Treatment inquiry log, WhatsApp appointment trigger'),
    ('Healthcare & Clinics', 'Pediatrics & Child Care Clinics', 'Pediatric Care & Baby Immunization Portal', 'Baby vaccination milestone schedule tracker, Pediatrician slot booking, Growth chart calculator, WhatsApp vaccine alert trigger', 'Vaccine schedule API, Age-based immunization calculator, Appointment logger'),
    ('Healthcare & Clinics', 'Ophthalmology & Eye Care Clinics', 'Eye Care & Optical Storefront Portal', 'Vision test appointment booking, Cataract/Lasik surgery information guide, Doctor schedule, Spectacle frames & optical catalog', 'Vision check slot booking, Optical product catalog API, Consultation queue'),
    ('Healthcare & Clinics', 'Orthopedics & Joint Care Clinics', 'Orthopedic & Joint Pain Clinic Portal', 'Joint pain self-assessment quiz, Orthopedic surgeon booking, Physiotherapy slot selector, Post-op rehabilitation video guides', 'Assessment score calculator, Consultation slot engine, WhatsApp booking notifier'),
    ('Healthcare & Clinics', 'Diabetology & Endocrinology Centers', 'Diabetology & Chronic Lifestyle Care Portal', 'Fasting blood sugar & HbA1c self-logger, Dietician consultation booking, Diabetic foot care clinic scheduler, Teleconsultation CTA', 'Patient sugar logs mock API, Diet plan downloads, Doctor appointment queue'),
    ('Healthcare & Clinics', 'ENT & Hearing Care Clinics', 'ENT & Audiology Assessment Portal', 'Hearing checkup booking, Sinus & allergy consultation scheduler, Ear endoscopy information, Digital hearing aids showcase', 'Audiogram test booking API, Doctor schedule manager, WhatsApp confirmation'),
    ('Healthcare & Clinics', 'Ayurveda & Homeopathy Wellness Centers', 'Holistic Healing & Natural Wellness Portal', 'Prakriti / Body constitution quiz, Panchakarma rejuvenation therapy booking, Herbal remedies catalog, Long-term case-taking scheduler', 'Holistic therapy booking API, Questionnaire response evaluator, Appointment queue'),
    ('Healthcare & Clinics', 'Multi-Specialty & General Outpatient Clinics', 'Multi-Specialty Polyclinic Management Portal', 'Multi-doctor department filter (Cardio, Gastro, Pulmo, General), Centralized token queue, Diagnostic test booking, Emergency CTA', 'Multi-department slot engine, Centralized receptionist token board, Patient history API'),
    ('Retail, Kirana & General Stores', 'Neighborhood Kirana & General Stores', 'Hyperlocal Digital Kirana PWA & WhatsApp Store', 'Daily staples (Rice, Dals, Oils, Spices), 2-column mobile card, inline quantity incrementer (- 1 +), 1-Click WhatsApp order, Free 30-min delivery banner', 'Product catalog state, Instant price updater, WhatsApp order generator, Digital Khata API'),
    ('Retail, Kirana & General Stores', 'Daily Needs Marts & Superettes', 'Neighborhood Mini-Mart Fast Order App', 'Packaged FMCG, Dairy & Cold Beverages, Snacks, Fast autocomplete search, Sticky floating cart, 1-Click WhatsApp delivery order', 'Product inventory search API, WhatsApp message compiler, Order logger'),
    ('Retail, Kirana & General Stores', 'Wholesale Grain Depots & Rice Merchants', 'B2B Wholesale Grain & Merchant Order Portal', 'Bulk bag pricing (25kg/50kg), Wholesale tiered discount calculator, Merchant credit khata, GST invoice generator, WhatsApp bulk cart', 'Tiered pricing API, Bulk cart formatter, Merchant credit ledger manager'),
    ('Agro Industry & Rice Mills', 'Modern Rice Mills & Commercial Hullers', 'B2B Industrial Miller Platform & Live Mandi Ticker', 'Sticky live mandi rate ticker (Raw/Steam Sona Masoori, Swarna 1061), Buhler Sortex Optical Sorter showcase, Bulk RFQ (Metric Tons) calculator, Freight dispatch log', 'Live Mandi rate ticker API, 1-click price controller, Bulk RFQ calculation engine, Lorry dispatch tracker'),
    ('Agro Industry & Rice Mills', 'Rice Processing & Value-Add Mills', 'Agro Commodity & By-Product Procurement Portal', 'Broken Rice (Nookalu) & Rice Bran (DORB) specs, Protein/moisture lab reports, Bulk MT inquiry form, Institutional buyer contracts', 'Commodity spec sheet API, Bulk RFQ calculator, Contract inquiry manager')
]

r_idx1 = 11
for broad, micro, app, fe, be in micro_specs_order:
    cnt = micro_counts[micro]
    pct = cnt / len(master_records)
    
    row_vals = [broad, micro, cnt, pct, app, fe, be]
    for c_idx, val in enumerate(row_vals, start=1):
        cell = ws1.cell(r_idx1, c_idx, val)
        cell.font = font_data
        cell.border = thin_border
        cell.fill = fill_zebra if r_idx1 % 2 == 0 else fill_white
        if c_idx in [1, 2]:
            cell.font = font_data_bold
        elif c_idx == 3:
            cell.alignment = align_center
            cell.number_format = '#,##0'
        elif c_idx == 4:
            cell.alignment = align_center
            cell.number_format = '0.0%'
        elif c_idx in [6, 7]:
            cell.alignment = align_wrap_left
        else:
            cell.alignment = align_left
    ws1.row_dimensions[r_idx1].height = 48
    r_idx1 += 1

# Total Row
tot_row = ['Total Commercial Pipeline', '14 Distinct Micro-Niches', 550, 1.0, 'Full Suite of Custom Solutions', 'Zero generic templates', 'Full interactive mock state']
for c_idx, val in enumerate(tot_row, start=1):
    cell = ws1.cell(r_idx1, c_idx, val)
    cell.font = font_data_bold
    cell.border = double_bottom
    cell.fill = PatternFill(start_color=navy_light, end_color=navy_light, fill_type='solid')
    if c_idx == 3:
        cell.alignment = align_center
        cell.number_format = '#,##0'
    elif c_idx == 4:
        cell.alignment = align_center
        cell.number_format = '0.0%'
    else:
        cell.alignment = align_left

# TABLE 2: District Cross-Tabulation
r_idx2 = r_idx1 + 2
ws1.cell(r_idx2, 1, '2. Micro-Business Distribution by District (Geographic Analysis)').font = font_section
r_idx2 += 1

tbl2_headers = ['Micro-Business Niche', 'Kadapa (YSR)', 'Krishna Dist', 'Vijayawada (NTR)', 'West Godavari', 'Total Leads']
for col_idx, h in enumerate(tbl2_headers, start=1):
    cell = ws1.cell(r_idx2, col_idx, h)
    cell.font = font_tbl_header
    cell.fill = fill_navy_header
    cell.alignment = align_center if col_idx > 1 else align_left
    cell.border = thin_border

# Compute cross-tab
cross_tab = defaultdict(lambda: defaultdict(int))
for r in master_records:
    cross_tab[r['micro_type']][r['district']] += 1

r_idx2 += 1
for _, micro, _, _, _ in micro_specs_order:
    kdp = cross_tab[micro]['Kadapa (YSR Kadapa)']
    kri = cross_tab[micro]['Krishna']
    vja = cross_tab[micro]['Vijayawada (NTR)']
    wg = cross_tab[micro]['West Godavari']
    tot = kdp + kri + vja + wg
    
    row_vals = [micro, kdp, kri, vja, wg, tot]
    for c_idx, val in enumerate(row_vals, start=1):
        cell = ws1.cell(r_idx2, c_idx, val)
        cell.font = font_data
        cell.border = thin_border
        cell.fill = fill_zebra if r_idx2 % 2 == 0 else fill_white
        if c_idx == 1:
            cell.font = font_data_bold
            cell.alignment = align_left
        else:
            cell.alignment = align_center
            cell.number_format = '#,##0'
    r_idx2 += 1

# Total Row for cross-tab
tot_ct = ['Total Leads per District', 300, 100, 50, 100, 550]
for c_idx, val in enumerate(tot_ct, start=1):
    cell = ws1.cell(r_idx2, c_idx, val)
    cell.font = font_data_bold
    cell.border = double_bottom
    cell.fill = PatternFill(start_color=navy_light, end_color=navy_light, fill_type='solid')
    if c_idx == 1:
        cell.alignment = align_left
    else:
        cell.alignment = align_center
        cell.number_format = '#,##0'

# ==========================================================
# SHEET 2: All Leads (Micro-Categorized Master)
# ==========================================================
ws2 = wb_new.create_sheet('All Leads (Micro-Categorized)')
ws2.views.sheetView[0].showGridLines = True

headers_master = [
    'Lead ID',
    'Primary Sector',
    'Micro-Business Niche',
    'Business Name',
    'Owner / Decision Maker',
    'Designation / Role',
    'Email Address',
    'Email Status',
    'Phone Number',
    'City / Town',
    'District',
    'Sub-Industry / Original Tag',
    'Tailored Web Application Archetype',
    'Key Frontend Demo Features',
    'Backend Demo Requirements',
    'GPS Coordinates',
    'Source File'
]

ws2.row_dimensions[1].height = 28
for c_idx, h in enumerate(headers_master, start=1):
    cell = ws2.cell(1, c_idx, h)
    cell.font = font_tbl_header
    cell.fill = fill_navy_header
    cell.border = thin_border
    cell.alignment = align_center if c_idx in [1, 8, 9, 16] else align_left

for r_idx, lead in enumerate(master_records, start=2):
    ws2.row_dimensions[r_idx].height = 20
    is_even = (r_idx % 2 == 0)
    base_fill = fill_zebra if is_even else fill_white
    
    vals = [
        lead['lead_id'],
        lead['broad_cat'],
        lead['micro_type'],
        lead['bname'],
        lead['owner'],
        lead['role'],
        lead['email'],
        lead['status'],
        lead['phone'],
        lead['city'],
        lead['district'],
        lead['sub_ind'],
        lead['app_name'],
        lead['fe_features'],
        lead['be_reqs'],
        lead['gps'],
        lead['source_file']
    ]
    
    for c_idx, val in enumerate(vals, start=1):
        cell = ws2.cell(r_idx, c_idx, val)
        cell.font = font_data
        cell.border = thin_border
        cell.fill = base_fill
        
        if c_idx == 1:
            cell.alignment = align_center
            cell.font = font_data_bold
        elif c_idx == 8: # Email Status
            cell.alignment = align_center
            if val == 'Verified':
                cell.fill = fill_green
                cell.font = Font(name='Segoe UI', size=9, bold=True, color=green_text)
            else:
                cell.fill = fill_amber
                cell.font = Font(name='Segoe UI', size=9, bold=True, color=amber_text)
        elif c_idx in [9, 16]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

ws2.freeze_panes = 'D2'
ws2.auto_filter.ref = f'A1:Q{len(master_records)+1}'

# ==========================================================
# FILTERED MICRO SHEETS (Sheets 3 to 7)
# ==========================================================
sub_sheets = [
    ('Dental Clinics (120)', lambda r: r['micro_type'] == 'Dental Clinics & Implant Centers'),
    ('Medical & Specialty (80)', lambda r: r['broad_cat'] == 'Healthcare & Clinics' and r['micro_type'] != 'Dental Clinics & Implant Centers'),
    ('Kirana & Retail Stores (178)', lambda r: r['micro_type'] in ['Neighborhood Kirana & General Stores', 'Daily Needs Marts & Superettes']),
    ('Wholesale Grain Depots (22)', lambda r: r['micro_type'] == 'Wholesale Grain Depots & Rice Merchants'),
    ('Modern Rice Mills (150)', lambda r: r['broad_cat'] == 'Agro Industry & Rice Mills')
]

for s_name, filter_fn in sub_sheets:
    ws_sub = wb_new.create_sheet(s_name)
    ws_sub.views.sheetView[0].showGridLines = True
    
    ws_sub.row_dimensions[1].height = 28
    for c_idx, h in enumerate(headers_master, start=1):
        cell = ws_sub.cell(1, c_idx, h)
        cell.font = font_tbl_header
        cell.fill = fill_navy_header
        cell.border = thin_border
        cell.alignment = align_center if c_idx in [1, 8, 9, 16] else align_left
        
    filtered_data = [r for r in master_records if filter_fn(r)]
    for r_idx, lead in enumerate(filtered_data, start=2):
        ws_sub.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)
        base_fill = fill_zebra if is_even else fill_white
        
        vals = [
            lead['lead_id'],
            lead['broad_cat'],
            lead['micro_type'],
            lead['bname'],
            lead['owner'],
            lead['role'],
            lead['email'],
            lead['status'],
            lead['phone'],
            lead['city'],
            lead['district'],
            lead['sub_ind'],
            lead['app_name'],
            lead['fe_features'],
            lead['be_reqs'],
            lead['gps'],
            lead['source_file']
        ]
        
        for c_idx, val in enumerate(vals, start=1):
            cell = ws_sub.cell(r_idx, c_idx, val)
            cell.font = font_data
            cell.border = thin_border
            cell.fill = base_fill
            
            if c_idx == 1:
                cell.alignment = align_center
                cell.font = font_data_bold
            elif c_idx == 8:
                cell.alignment = align_center
                if val == 'Verified':
                    cell.fill = fill_green
                    cell.font = Font(name='Segoe UI', size=9, bold=True, color=green_text)
                else:
                    cell.fill = fill_amber
                    cell.font = Font(name='Segoe UI', size=9, bold=True, color=amber_text)
            elif c_idx in [9, 16]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
    ws_sub.freeze_panes = 'D2'
    ws_sub.auto_filter.ref = f'A1:Q{len(filtered_data)+1}'

# Column widths auto-adjustment
for ws in wb_new.worksheets:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col[:40]:
            v_str = str(cell.value or '')
            if '\n' in v_str:
                l_len = max(len(line) for line in v_str.split('\n'))
            else:
                l_len = len(v_str)
            if l_len > max_len: max_len = l_len
        if ws.title == 'Micro Analytics & Matrix':
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 15), 45)
        else:
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

output_file = 'Micro_Categorized_Leads_and_Web_Solutions_Directory.xlsx'
wb_new.save(output_file)
print(f'Successfully generated NEW Excel workbook: {output_file}')
